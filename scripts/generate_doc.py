"""Generate Music2txt project documentation as a Word (.docx) file."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "Music2txt_專案說明.docx"


def _set_run_font(run, font_name: str = "微軟正黑體", size_pt: float | None = None) -> None:
    run.font.name = font_name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)
    if size_pt is not None:
        run.font.size = Pt(size_pt)


def _add_heading(doc: Document, text: str, level: int = 1) -> None:
    heading = doc.add_heading(text, level=level)
    for run in heading.runs:
        _set_run_font(run, size_pt=16 if level == 1 else 13)


def _add_para(doc: Document, text: str, *, bold: bool = False) -> None:
    para = doc.add_paragraph()
    run = para.add_run(text)
    run.bold = bold
    _set_run_font(run, size_pt=11)


def _add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        para = doc.add_paragraph(style="List Bullet")
        run = para.add_run(item)
        _set_run_font(run, size_pt=11)


def _add_numbered(doc: Document, items: list[str]) -> None:
    for item in items:
        para = doc.add_paragraph(style="List Number")
        run = para.add_run(item)
        _set_run_font(run, size_pt=11)


def _add_table(doc: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ""
        run = cell.paragraphs[0].add_run(header)
        run.bold = True
        _set_run_font(run, size_pt=11)
    for r_idx, row in enumerate(rows, start=1):
        for c_idx, value in enumerate(row):
            cell = table.rows[r_idx].cells[c_idx]
            cell.text = ""
            run = cell.paragraphs[0].add_run(value)
            _set_run_font(run, size_pt=11)
    doc.add_paragraph()


def build_document() -> Document:
    doc = Document()

    style = doc.styles["Normal"]
    style.font.name = "微軟正黑體"
    style.font.size = Pt(11)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "微軟正黑體")

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Music2txt 專案說明文件")
    run.bold = True
    _set_run_font(run, size_pt=20)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = sub.add_run("離線 MP3／音訊轉文字小工具")
    _set_run_font(run, size_pt=12)

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = meta.add_run("文件版本：初版說明｜適用專案路徑：Music2txt_project")
    _set_run_font(run, size_pt=10)

    # 1. 專案概述
    _add_heading(doc, "一、專案概述", 1)
    _add_para(
        doc,
        "Music2txt 是一套以 Python 撰寫的離線語音轉文字桌面小工具。"
        "使用者可透過圖形介面開啟音檔，選擇處理模式後進行轉寫，"
        "結果會顯示於文字框，並可複製或另存為 .txt 檔。",
    )
    _add_para(doc, "適用場景：", bold=True)
    _add_bullets(
        doc,
        [
            "會議講話錄音：訪談、會議、podcast 等一般語音內容。",
            "歌曲（帶伴奏）：先分離人聲再轉寫，改善伴奏干擾造成的誤聽。",
            "中英混合句型：使用 Whisper 自動語系偵測，可處理中英文夾雜。",
        ],
    )
    _add_para(
        doc,
        "設計重點：轉寫過程完全在本機執行，不需雲端 API Key；"
        "僅在首次下載模型時需要網路，之後可離線使用。",
    )

    # 2. 系統需求
    _add_heading(doc, "二、系統需求", 1)
    _add_table(
        doc,
        ["項目", "說明"],
        [
            ["作業系統", "Windows 10／11（本專案開發與驗證環境）"],
            ["Python", "3.10（專案使用虛擬環境資料夾 env）"],
            ["硬體建議", "建議 8GB 以上記憶體；歌曲模式較吃 CPU／RAM"],
            ["磁碟空間", "模型與套件約需數 GB（Whisper small、Demucs 等）"],
            ["FFmpeg", "建議安裝並加入 PATH，作為部分音訊格式讀取備援"],
            ["網路", "首次執行需下載模型；下載完成後可完全離線"],
        ],
    )

    # 3. 使用技術與角色
    _add_heading(doc, "三、使用技術與角色", 1)
    _add_table(
        doc,
        ["技術／元件", "角色"],
        [
            ["Python 3.10 + venv（env）", "執行環境與依賴隔離"],
            ["tkinter／ttk", "圖形介面：開檔、模式選擇、進度條、文字框、複製／另存"],
            ["threading + queue", "背景執行轉檔，避免 UI 凍結；進度與結果回傳主執行緒"],
            ["faster-whisper", "語音轉文字引擎（預設模型 small，CPU／int8，自動語系）"],
            ["demucs（htdemucs）", "歌曲模式：從混音中分離出人聲音軌"],
            ["torch／torchaudio", "Demucs 推論所需的深度學習運算框架"],
            ["soundfile", "音訊檔讀寫輔助"],
            ["FFmpeg（系統）", "部分格式（如部分 MP3）讀取失敗時的備援解碼"],
        ],
    )

    # 4. 技術如何串接
    _add_heading(doc, "四、技術如何串接（處理流程）", 1)
    _add_para(doc, "程式啟動路徑：", bold=True)
    _add_para(doc, "main.py → app/ui.py（建立視窗與事件）→ 背景工作呼叫轉寫／人聲分離模組。")

    _add_heading(doc, "4.1 會議講話模式", 2)
    _add_numbered(
        doc,
        [
            "使用者選擇音檔並按「開始轉檔」。",
            "背景執行緒載入 faster-whisper 模型（small）。",
            "直接對音檔轉寫（language=None，自動偵測語系，含中英混合）。",
            "進度條更新狀態；完成後文字寫入文字框。",
        ],
    )
    _add_para(doc, "進度大致分配：載入模型約 20% + 轉寫段落約 80%。")

    _add_heading(doc, "4.2 歌曲（帶伴奏）模式", 2)
    _add_numbered(
        doc,
        [
            "使用者選擇音檔並選擇「歌曲（帶伴奏）」後按「開始轉檔」。",
            "demucs 載入 htdemucs，分離出 vocals 人聲並暫存為 wav。",
            "faster-whisper 對人聲檔轉寫。",
            "結果顯示於文字框；暫存人聲檔會自動清理。",
        ],
    )
    _add_para(doc, "進度大致分配：人聲分離約 40% +（載入模型與轉寫）約 60%。")

    _add_heading(doc, "4.3 模組對應", 2)
    _add_table(
        doc,
        ["檔案", "職責"],
        [
            ["main.py", "程式入口，啟動 UI"],
            ["app/ui.py", "tkinter 介面、執行緒、進度與結果更新、複製／另存"],
            ["app/transcriber.py", "封裝 faster-whisper 載入與轉寫、進度回報"],
            ["app/vocal_separator.py", "封裝 Demucs 人聲分離與暫存清理"],
            ["requirements.txt", "第三方套件依賴清單"],
        ],
    )

    # 5. 依賴項列表
    _add_heading(doc, "五、依賴項列表", 1)

    _add_heading(doc, "5.1 Python 第三方套件（requirements.txt）", 2)
    _add_table(
        doc,
        ["套件", "用途"],
        [
            ["faster-whisper>=1.1.0", "離線語音轉文字（Whisper 加速版）"],
            ["demucs>=4.0.1", "歌曲人聲／伴奏分離"],
            ["torch>=2.0.0", "深度學習運算（Demucs 依賴）"],
            ["torchaudio>=2.0.0", "音訊相關支援（Demucs 生態）"],
            ["soundfile>=0.12.1", "音訊檔讀寫"],
        ],
    )
    _add_para(
        doc,
        "安裝上述套件時，pip 會一併帶入相關間接依賴"
        "（例如 ctranslate2、av、numpy、huggingface-hub 等），無需手動逐一安裝。",
    )

    _add_heading(doc, "5.2 Python 標準庫（無需 pip 安裝）", 2)
    _add_bullets(
        doc,
        [
            "tkinter／ttk／filedialog／messagebox：介面與對話框",
            "threading、queue：背景工作與訊息佇列",
            "pathlib、tempfile、shutil：路徑與暫存檔管理",
            "traceback：錯誤細節輸出（方便除錯）",
        ],
    )

    _add_heading(doc, "5.3 系統層依賴", 2)
    _add_bullets(
        doc,
        [
            "FFmpeg：建議安裝並加入系統 PATH。Demucs 在部分格式讀取失敗時會嘗試以 FFmpeg 備援。",
            "虛擬環境資料夾 env：專案 Python 套件安裝位置，勿刪除後直接執行（需重新建立並安裝依賴）。",
        ],
    )

    # 6. 環境安裝與操作說明
    _add_heading(doc, "六、環境安裝與操作說明", 1)

    _add_heading(doc, "6.1 建立虛擬環境與安裝依賴（若尚未建立）", 2)
    _add_para(doc, "在專案根目錄 PowerShell 執行：")
    _add_para(doc, "python -m venv env")
    _add_para(doc, ".\\env\\Scripts\\Activate.ps1")
    _add_para(doc, "pip install -r requirements.txt")

    _add_heading(doc, "6.2 啟動程式", 2)
    _add_para(doc, "方式 A（建議，不需先 Activate）：")
    _add_para(doc, ".\\env\\Scripts\\python.exe main.py")
    _add_para(doc, "方式 B（已 Activate 虛擬環境時）：")
    _add_para(doc, "python main.py")

    _add_heading(doc, "6.3 介面操作步驟", 2)
    _add_numbered(
        doc,
        [
            "按「開啟檔案」，選擇 .mp3／.wav／.m4a 等音訊檔。",
            "選擇模式：「會議講話」或「歌曲（帶伴奏）」。",
            "按「開始轉檔」。轉檔中按鈕會停用，進度條與狀態列會更新。",
            "完成後，轉寫文字出現在下方文字框（可直接編輯）。",
            "可按「複製全文」或「另存為 .txt」匯出結果。",
        ],
    )

    _add_heading(doc, "6.4 重新產生本說明文件", 2)
    _add_para(doc, "若需更新文件內容後重新產出 Word：")
    _add_para(doc, ".\\env\\Scripts\\python.exe scripts\\generate_doc.py")
    _add_para(doc, "產出檔案：專案根目錄的 Music2txt_專案說明.docx")

    # 7. 專案結構
    _add_heading(doc, "七、專案結構", 1)
    _add_para(doc, "主要目錄與檔案如下（精簡）：")
    structure_lines = [
        "Music2txt_project/",
        "├── main.py                 # 程式入口",
        "├── requirements.txt        # 第三方依賴",
        "├── Music2txt_專案說明.docx  # 本說明文件",
        "├── app/",
        "│   ├── __init__.py",
        "│   ├── ui.py               # 圖形介面與工作排程",
        "│   ├── transcriber.py      # Whisper 轉寫",
        "│   └── vocal_separator.py  # Demucs 人聲分離",
        "├── scripts/",
        "│   └── generate_doc.py     # 產生本 Word 的腳本",
        "├── backup/                 # 程式碼備份",
        "├── env/                    # Python 虛擬環境（勿提交敏感環境時可忽略）",
        "└── .gitignore",
    ]
    for line in structure_lines:
        para = doc.add_paragraph()
        run = para.add_run(line)
        _set_run_font(run, font_name="Consolas", size_pt=10)
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "微軟正黑體")

    # 8. 已知限制
    _add_heading(doc, "八、已知限制與注意事項", 1)
    _add_bullets(
        doc,
        [
            "歌曲轉歌詞無法保證百分之百正確；和聲、效果器、咬字不清、強伴奏仍可能造成漏字或誤聽。",
            "首次執行會下載 Whisper／Demucs 模型，需一次網路與較長等待；之後可離線。",
            "CPU 模式轉寫與人聲分離較耗時，長音檔請耐心等候。",
            "轉檔中請勿重複操作；介面已暫時停用「開始轉檔」以避免重複觸發。",
            "Windows 若未啟用開發人員模式，Hugging Face 快取可能出現 symlink 警告，不影響功能，但可能多佔磁碟空間。",
        ],
    )

    # 9. 常見問題
    _add_heading(doc, "九、常見問題排解", 1)
    _add_table(
        doc,
        ["狀況", "可能原因與處理"],
        [
            [
                "讀取 MP3 失敗／轉檔報錯提及 FFmpeg",
                "請安裝 FFmpeg 並將其加入系統 PATH，重新開啟終端機後再試。",
            ],
            [
                "首次很慢或顯示下載相關警告",
                "正在下載模型；請保持網路暢通。完成後同一模型不會重複下載。",
            ],
            [
                "文字框顯示「沒有辨識到語音內容」",
                "音檔可能過短、靜音、人聲極弱，或歌曲模式分離後仍難辨識。可改試另一段音檔或確認模式選擇。",
            ],
            [
                "按開始沒反應或找不到 python",
                "請使用 .\\env\\Scripts\\python.exe main.py，確認在專案根目錄執行。",
            ],
            [
                "記憶體不足或程式卡住很久",
                "歌曲模式較吃資源；可先用較短音檔測試，或關閉其他大型程式。",
            ],
        ],
    )

    _add_para(
        doc,
        "若問題持續，請保留錯誤訊息內容（程式失敗時會彈出對話框，細節也可能印在啟動用的終端機），以便進一步排查。",
    )

    return doc


def main() -> None:
    doc = build_document()
    doc.save(OUTPUT)
    print(f"已產生：{OUTPUT}")


if __name__ == "__main__":
    main()
